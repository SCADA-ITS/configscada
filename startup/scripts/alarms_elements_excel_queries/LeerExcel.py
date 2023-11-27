import openpyxl as xl
import os
import shutil
from translation import todatabase, todatabasetypes 
import sys
import psycopg2


#INFO DATABASE

# Variables para la conexion a BBDD
HOST = sys.argv[1]
PORT = sys.argv[2]
DB = "rits"
USER = "rits"
PASS = "rits"

#Function that returns element_id given the alias 
def get_element_id_by_sheet_name(name):
    try:
        #DB Connection
        connectionChain = "host=%s port=%s user=%s password=%s dbname=%s" % (
            HOST, PORT, USER, PASS, DB)
        connection = psycopg2.connect(connectionChain)
        # Cursor to operate on the DB
        cur = connection.cursor()

        cur.execute("SELECT element_type_id FROM master.element_types WHERE label_alias = \'" + name + '\';')
        info = cur.fetchone()
        cur.close()
        connection.close()
        #Get the result and return the first value
        if (info == None):
            return name
        else:
            return info[0]

    except AttributeError:
        print(ValueError)

#Given an string (update queries) execute and commit them to the DB
def update_database_table (updates):
    try:
        #DB Connection
        connectionChain = "host=%s port=%s user=%s password=%s dbname=%s" % (
            HOST, PORT, USER, PASS, DB)
        connection = psycopg2.connect(connectionChain)
        # Cursor to operate on the DB
        cur = connection.cursor()
        cur.execute(updates)
        connection.commit()
        cur.close()
        connection.close()

    except AttributeError:
        print(ValueError)

def getaliasbyid(element_type_id):
    try:
        #DB Connection
        connectionChain = "host=%s port=%s user=%s password=%s dbname=%s" % (
            HOST, PORT, USER, PASS, DB)
        connection = psycopg2.connect(connectionChain)
        # Cursor to operate on the DB
        cur = connection.cursor()
        consult = "SELECT alias FROM master.element_types WHERE element_type_id = "+ str(element_type_id) +";"
        cur.execute(consult)

        info = cur.fetchone()

        cur.close()
        connection.close()
        if info != None:
            info = info[0]
        return info

    except AttributeError:
        print(ValueError)
#Given the the excel file's path
#If option = 1 --> execute updates on database
#if option = 2 --> create .sql file with all update consults
def read_excel(path, option):
    # To open the workbook
    excel_file = xl.load_workbook(path)
    
    #If the update folder exists, erase all the data from it
    if os.path.exists('UPDATES/'):
        shutil.rmtree('UPDATES/')
    os.mkdir('UPDATES/')

    #Number of sheets in the excel file
    number_sheets = len(excel_file.sheetnames)
    #All the sheets
    sheets = excel_file.sheetnames

    #Readying sheets loop
    for i in range(0, number_sheets):
        actual_sheet_name = excel_file[sheets[i]]
        excel_file.active =  actual_sheet_name
        active_sheet = excel_file.active

        max_rows = active_sheet.max_row +1
        element_type_id = int(actual_sheet_name.title.split('.-')[0])


        #Loop for generating all the update queries
        for j in range (2, max_rows):
            #Get all the variables for the query
            alarm_config_id = active_sheet.cell(row = j, column = 1).value
            label_alias = active_sheet.cell(row = j, column = 2).value
            
            label_alias_db = todatabase(label_alias)
            alarm_level_id = int(active_sheet.cell(row = j, column = 3).value) +2
            enabled = active_sheet.cell(row = j, column = 4).value
            if enabled == '=TRUE()':
                enabled = 'True'
            elif enabled == '=FALSE()':
                enabled = 'False'
            treatment = 'True'
            if alarm_level_id == 2:
                treatment = 'False'
            if enabled != 'False':
                if (label_alias_db != None):
                    consult = 'UPDATE conf.alarm_configs SET alarm_level_id = ' + str(alarm_level_id) + ', enabled = ' + str(enabled) + ', treatment = ' +str(treatment) + ' WHERE element_type_id = ' + str(element_type_id) + ' AND alarm_config_id = ' + str(alarm_config_id) + '; --'+ label_alias_db+'\n'        
                    #If option 2, write the consult on the sql file
                    if (option == 2):
                        name_file = getaliasbyid(element_type_id)
                        name_file = str(element_type_id) +'.' + name_file.lower()
                        name_file = name_file.replace(' ', '_')
                        file_name = 'UPDATES/' + name_file +'.sql'
                        f = open(file_name , 'a')
                        f.write(consult)
                        f.close()
                    #If option 1, call the function that executes the update query
                    elif(option == 1):
                        update_database_table(consult)
